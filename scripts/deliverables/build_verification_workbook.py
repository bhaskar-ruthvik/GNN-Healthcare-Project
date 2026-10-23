"""Verification workbook for the professors: every reported number with its source file, and formula checks that recompute the
headline claims inside the workbook. AGGREGATES ONLY: no referral-, patient- or provider-level rows are read into or written to it.
Output: deliverables/referral_gnn_results_verification.xlsx (refuses to overwrite; pass --out for a new name).
Run from anywhere (data under the workspace, see referral_gnn/paths.py):  python3 scripts/deliverables/build_verification_workbook.py [--out deliverables/<name>.xlsx]"""
import os, sys, pathlib  # bootstrap: make referral_gnn importable from any scripts/ subfolder (or `pip install -e .`)
sys.path.insert(0, str(next(p for p in pathlib.Path(__file__).resolve().parents if (p / "referral_gnn" / "paths.py").is_file())))
from referral_gnn.paths import workspace, script_path, code_file  # noqa: E402,F401
import argparse, glob, json, os, re
from datetime import date
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

ap = argparse.ArgumentParser(); ap.add_argument("--out", default="deliverables/referral_gnn_results_verification.xlsx"); args = ap.parse_args()
root = str(workspace()); os.chdir(root)
if os.path.exists(args.out): raise SystemExit(f"REFUSED: {args.out} exists (use --out)")
A = "runs/analysis"; J = lambda p: json.load(open(p))
FONT = "Arial"; F = Font(name=FONT, size=10); FB = Font(name=FONT, size=10, bold=True); FH = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FT = Font(name=FONT, size=14, bold=True); FS = Font(name=FONT, size=9, italic=True, color="555555"); FBLUE = Font(name=FONT, size=10, color="0000FF")
HDR = PatternFill("solid", fgColor="44546A"); CHK = PatternFill("solid", fgColor="E2EFDA"); KEY = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
wb = Workbook(); REG = {}          # REG[(sheet, key)] = absolute cell reference used by the Claims sheet

def fmt_for(col):
    c = col.lower()
    if re.search(r"(^|_)(n|n_seeds|referrals|patients|pcps|count|rank|seed|epoch|best_epoch|epochs_run|distinct|n_points|n_features|quintile|q|k)$", c): return "0"
    if re.search(r"p_holm|^p$|p_value|_p$|p_one|ari|gini|tau|r2|beta|spearman|ratio|selection|mrr|^rr$|volpct", c): return "0.000"
    if re.search(r"gain|mips|quality|excess|dist|travel|_mi$|miles|effect_pts|coef", c): return "0.00"
    if re.search(r"fit|busy|load|new|r5|r10|share|rate|insurer|category|within|better|equal|worse|low_mips|none|seen|conc|early|lowexp|nonmetro|good|agree|overlap|avail|actual|familiar|headroom|^ins$|^cat$|^busy$|^new$|fraction|frac", c): return "0.0%"
    return None

def sheet(name, title, note, src):
    ws = wb.create_sheet(name); ws["A1"] = title; ws["A1"].font = FT
    ws["A2"] = note; ws["A2"].font = FS; ws["A3"] = "Source: " + (src if isinstance(src, str) else "; ".join(src)); ws["A3"].font = FS
    ws.sheet_view.showGridLines = False; return ws

def table(ws, df, r0, c0=1, name=None, fmts=None, widths=True):
    """Write a DataFrame as a styled table at (r0, c0); returns {column: (first_row, last_row, col_letter)}."""
    fmts = fmts or {}
    for j, col in enumerate(df.columns):
        cell = ws.cell(row=r0, column=c0 + j, value=str(col)); cell.font = FH; cell.fill = HDR; cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(df.itertuples(index=False)):
        for j, v in enumerate(row):
            if isinstance(v, (np.floating, float)) and not np.isfinite(v): v = None
            elif isinstance(v, np.generic): v = v.item()
            elif isinstance(v, (list, dict)): v = json.dumps(v)
            cell = ws.cell(row=r0 + 1 + i, column=c0 + j, value=v); cell.font = F; cell.border = THIN
            f = fmts.get(df.columns[j], fmt_for(str(df.columns[j])))
            if f and isinstance(v, (int, float)) and not isinstance(v, bool): cell.number_format = f
    if widths:
        for j, col in enumerate(df.columns):
            w = max([len(str(col))] + [len(str(x)) for x in df.iloc[:, j].head(200)]) if len(df) else len(str(col))
            ws.column_dimensions[L(c0 + j)].width = max(ws.column_dimensions[L(c0 + j)].width or 0, min(max(9, w * 0.9 + 2), 60))
    if name and len(df):
        t = Table(displayName=re.sub(r"\W", "_", name), ref=f"{L(c0)}{r0}:{L(c0 + len(df.columns) - 1)}{r0 + len(df)}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True); ws.add_table(t)
    return {col: (r0 + 1, r0 + len(df), L(c0 + j)) for j, col in enumerate(df.columns)}

def cell_of(ws, cols, col, i): r0, _, c = cols[col]; return f"'{ws.title}'!${c}${r0 + i}"
def ci_split(df, col, lo="ci_lo", hi="ci_hi"):
    df = df.copy(); df[lo] = [x[0] if isinstance(x, (list, tuple)) else None for x in df[col]]; df[hi] = [x[1] if isinstance(x, (list, tuple)) else None for x in df[col]]
    return df.drop(columns=[col])
def checkcol(ws, cols, rows, formula_of, header, col_letter, fmt=None):
    """Add a formula column next to a table (e.g. a recomputed difference)."""
    r0 = list(cols.values())[0][0] - 1; ws[f"{col_letter}{r0}"] = header; ws[f"{col_letter}{r0}"].font = FH; ws[f"{col_letter}{r0}"].fill = PatternFill("solid", fgColor="548235")
    ws.column_dimensions[col_letter].width = max(14, len(header) * 0.9)
    for i in range(rows):
        c = ws[f"{col_letter}{r0 + 1 + i}"]; c.value = formula_of(r0 + 1 + i); c.font = F; c.fill = CHK
        if fmt: c.number_format = fmt

# =============================================================== plain-language names ===========================================================
# Readers see what each experiment and model IS, never the project's internal step / hypothesis / cell codes.
PLAIN_CLAIM = {"H1a": "Message passing recovers history", "H1b": "Learned beats fixed summaries", "H2_low_history": "Graph helps low-history PCPs",
               "H2_thin_network": "Graph helps thin-network PCPs", "H3": "Insurer-compatible new picks"}
PLAIN_CELL = {"graph": "full GNN", "ngp": "no message passing (with PCP)", "ngp_nohist": "no message passing without history", "gnn_nohist": "GNN without history",
              "table": "pure table", "flat_rf": "table + graph summaries", "mlp_rf": "MLP on graph summaries", "info_nohist": "richer table without history",
              "info": "richer table", "mlp_info": "MLP on the richer table", "insmatch": "insurer-share control", "g": "no message passing, no PCP, no history"}
PLAIN_SLICE = {"low_history": "low-history PCPs", "thin_network": "thin-network PCPs", "rare": "rare conditions"}
PLAIN_FAM = {"G": "GNN", "T": "pure table", "F": "table + graph summaries", "B_GT": "GNN + table (earlier steered blend)",
             "B_GF": "GNN + table with graph summaries", "B_GTF": "three-way: GNN + table + table with graph summaries"}
def plain_component(n):
    """Readable name of a graph-value comparison (the result files use internal cell names)."""
    fixed = {"prerequisite: Graph − NG+P-nohist": "prerequisite: full GNN − no message passing without history",
             "GNN-nohist − NG+P-nohist": "GNN − no message passing (both without history)", "joint half-gap": "recovers at least half of history's value"}
    if n in fixed: return fixed[n]
    m = re.match(r"^(absolute|interaction): Graph − (\w+) \((\w+)\)$", n)
    if m: return f"{m.group(1)}: full GNN − {PLAIN_CELL[m.group(2)]} ({PLAIN_SLICE[m.group(3)]})"
    m = re.match(r"^GNN-nohist − (\w+)( \(matched newness\))?$", n)
    if m: return f"GNN without history − {PLAIN_CELL[m.group(1)]}{m.group(2) or ''}"
    raise ValueError(n)

# =============================================================== README ===========================================================
ws = wb.active; ws.title = "README"; ws.sheet_view.showGridLines = False
ws["A1"] = "Referral GNN — results verification workbook"; ws["A1"].font = Font(name=FONT, size=16, bold=True)
readme = [
    ("Built", f"{date.today().isoformat()} by scripts/deliverables/build_verification_workbook.py. Every number is read from the result files named on each sheet (row 3); nothing is typed by hand."),
    ("Code", "https://github.com/bhaskar-ruthvik/GNN-Healthcare-Project — every result file named on the sheets is produced by the scripts there (REPRODUCE.md maps results to commands)."),
    ("Privacy", "Aggregates only. No patient-, referral- or provider-level values are in this workbook. The underlying data and per-referral outputs stay on the project machine."),
    ("How to verify", "1) Each sheet names its source file(s) in row 3; the file and the script that wrote it are in the code repository. "
                      "2) Green columns are formulas that recompute a reported number from the other cells (differences, adjusted p-values, minimums over a grid). "
                      "3) The 'Claims' sheet lists every headline statement in the decks with the cells that support it and a TRUE/FALSE check formula (all TRUE when built). "
                      "Columns labelled 'test criterion' (sheet S12a_Mechanism) are pre-registered test outcomes; FALSE there is a result, not an error."),
    ("Years", "2018 = validation year (all choices made here). 2019 = held-out year, used once per pre-stated check through a guard that requires a recorded decision (sheet 'Decisions')."),
    ("Cohort", "First-time referrals (the patient had never seen that specialist), ranked against eligible specialists: same specialty, active in the last 365 days, within 25 miles (filled to 5, or to K for top-K), not already seen by the patient. 5,550 referrals in 2018, 2,668 in 2019 (sheet 'Cohort')."),
    ("Lenses", "Fit = the specialist the PCP chose is in the top 5 (top K). Load = share of top picks going to the busiest 10% of specialists in that specialty (busy = above the specialty's 90th percentile of first-time referrals in the training years); 'top 5'/'slots' versions average over the list. "
               "New to the PCP = share of the list the PCP had never referred to before. MIPS gain = mean MIPS of the list minus MIPS of the chosen specialist (points). Travel = distance from the patient (miles)."),
    ("30% quality weight", "Unless stated, rankings are 0.7 × z(model score) + 0.3 × z(specialist MIPS as of the referral date), z-scored within each referral's candidate list ('q30' / 'blend0.3')."),
    ("Steering", "Adds a·(new to the PCP) − b·(in the busiest 10%) to the ranking score; both flags are known at referral time. The steered equal blend uses 0.5·z(GNN) + 0.5·z(table), a = 0.3, b = 0.4 (chosen on 2018)."),
    ("Intervals", "95% intervals resample whole PCPs (119 PCPs), unless the sheet says patient-cluster. 'Matched fit' comparisons dilute the more accurate model with random picks until fits match."),
    ("Seeds", "Models are averaged over seeds; 'pairing' = one GNN run with one table run (GNN s42 / s1 / s42_dup with table s42 / s1 / s7). The graph-value study and the blend with graph summaries pair by seed ID (42 / 1 / 7; a third GNN seed was trained for them)."),
    ("Graph value", "'What the graph adds': pre-registered tests against models with the same information but no message passing. 'Matched fit (post hoc)': descriptive. 'Blend with graph summaries': the steered recommender including the table with graph summaries (setting chosen on 2018, checked once on 2019)."),
    ("Formats", "Percentages are stored as fractions (0.403 displays as 40.3%). MIPS in points (0–100). Distances in miles."),
    ("MRR", "Mean reciprocal rank appears only in the run register as a training/selection sanity value; no result is judged on it."),
]
for i, (k, v) in enumerate(readme):
    ws.cell(row=3 + i, column=1, value=k).font = FB; c = ws.cell(row=3 + i, column=2, value=v); c.font = F; c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 140
SHEET_INDEX_ROW = 4 + len(readme)

# =============================================================== Cohort ===========================================================
ws = sheet("Cohort", "Cohort funnel: first-time referrals → primary evaluation cohort", "Counts per exclusion step (STEP10_PROTOCOL §2.1).", [f"{A}/step10_3/validation_summary.json", f"{A}/step10_3/test_summary.json"])
fv, ft = J(f"{A}/step10_3/validation_summary.json")["funnel"], J(f"{A}/step10_3/test_summary.json")["funnel"]
LAB = {"first_time": "First-time referrals in the year", "unknown_specialty": "− specialty unknown before the referral date", "non_choice_specialty": "− facility-based specialties (no real choice)",
       "non_specific_specialty": "− non-specific specialties", "referred_inactive": "− chosen specialist inactive in the prior year", "referred_mips_none": "− chosen specialist has no MIPS score",
       "eligible_lt5": "− fewer than 5 eligible specialists", "patient_zip_missing": "− patient ZIP missing", "primary": "Primary cohort"}
df = pd.DataFrame([{"step": LAB.get(k, k), "key": k, "2018": fv[k], "2019": ft.get(k)} for k in fv])
cols = table(ws, df, 5, name="cohort", fmts={"2018": "#,##0", "2019": "#,##0"})
n = len(df); checkcol(ws, cols, 1, lambda r: f"=C6-SUM(C7:C{5 + n - 1})", "2018: first-time − exclusions", "E", "#,##0"); ws["E6"].value = f"=C6-SUM(C7:C{4 + n})"
ws["F5"] = "matches primary?"; ws["F5"].font = FH; ws["F5"].fill = PatternFill("solid", fgColor="548235"); ws["F6"] = f"=E6=C{5 + n}"; ws["F6"].fill = CHK
REG[("Cohort", "n2018")] = f"'Cohort'!$C${5 + n}"; REG[("Cohort", "n2019")] = f"'Cohort'!$D${5 + n}"

# =============================================================== Runs ===========================================================
rows = []
for d in sorted(glob.glob("runs/2026*")):
    if not os.path.isdir(d) or "INCOMPLETE" in d or not os.path.exists(f"{d}/summary.json") or not os.path.exists(f"{d}/run_manifest.json"): continue
    m, s = J(f"{d}/run_manifest.json"), J(f"{d}/summary.json"); mc, tc = m.get("model_config", {}) or {}, m.get("train_config", {}) or {}
    flags = {k: v for k, v in {**mc, **tc}.items() if k in ("no_graph", "mask_demographics", "temporal_edges", "loss", "volume_mode", "no_insurer_history", "table_offset", "tab_dim",
                                                            "history_dropout", "scorer_skip", "exclude_fold", "n_folds", "train_end", "pos_weight_alpha") and v not in (None, False, 0, "", "softmax")}
    name = os.path.basename(d); kind = m.get("kind") or ("gnn" if m.get("schema", "").startswith("three-node") else "other")
    inp = m.get("inputs_path", ""); ver = os.path.basename(os.path.dirname(inp)) if inp else ""
    rows.append({"run_dir": name, "started": name[:15], "kind": kind, "seed": m.get("seed"), "inputs": ver, "inputs_sha256": (m.get("inputs_sha256") or "")[:12],
                 "code_sha256": (m.get("code_sha256") or "")[:12], "settings": json.dumps(flags) if flags else "", "params": json.dumps(m.get("params"))[:120] if m.get("params") else "",
                 "best_epoch": s.get("best_epoch"), "epochs_run": s.get("epochs_run"), "selection_metric": s.get("selection_metric"), "selection_value (sanity only)": s.get("best_val_selection")})
ws = sheet("Runs", "Run register: every completed training run", "One row per run directory with a summary.json (incomplete runs excluded). Settings show only non-default flags; hashes truncated to 12 characters.", "runs/*/run_manifest.json, runs/*/summary.json")
table(ws, pd.DataFrame(rows), 5, name="runs", fmts={"selection_value (sanity only)": "0.000"})

# =============================================================== Decisions ===========================================================
dec = J(f"{A}/step10_decisions.json")["entries"]
ws = sheet("Decisions", "Held-out year guard: decisions recorded before each 2019 pass", "The 2019 evaluation refuses to run unless an entry exists for the same step and runs, with the SHA-256 of every 2018 output file unchanged.", f"{A}/step10_decisions.json")
table(ws, pd.DataFrame([{"step": e["step"], "recorded_utc": e.get("recorded_utc"), "n_runs": len(e["runs"]), "final_or_primary": e.get("final_config_10_4_to_10_6") or e.get("primary_contrast_config"),
                         "amendment": e.get("amendment"), "n_2018_files_hashed": len(e.get("validation_files_sha256", {})), "note": (e.get("note") or "")[:200]} for e in dec]), 5, name="decisions")

# =============================================================== Step 10 arms ===========================================================
KEEP = ["arm", "n", "n_seeds", "r5", "gain5", "gain5_lo", "gain5_hi", "gain5_cons", "gain5_cons_lo", "gain5_cons_hi", "gain1", "busy1", "busy1_lo", "busy1_hi", "busy5", "new_pcp5",
        "dist5_median", "dist1_median", "within1_1", "good_new5", "insurer5", "gini_within", "r5_lo", "r5_hi"]
def arms_sheet(name, step_dir, title, note):
    for split, yr in (("validation", "2018"), ("test", "2019")):
        p = f"{A}/{step_dir}/{split}_arms.csv"
        if not os.path.exists(p): continue
        d = pd.read_csv(p); d = d[[c for c in KEEP if c in d.columns]]
        ws = sheet(f"{name}_{yr}", f"{title} — {yr}", note + " Columns: see 'Columns'. _lo/_hi = 95% interval.", p)
        cols = table(ws, d, 5, name=f"{name}_{yr}")
        for i, a in enumerate(d.arm): REG[(f"{name}_{yr}", a)] = i
        REG[(f"{name}_{yr}", "_cols")] = cols
arms_sheet("S10_Final", "step10_3", "Step 10.3: final model (v5) vs heuristics, all re-rankings", "Every scorer × re-ranking (base, re-sort top K by MIPS, blend with λ MIPS weight). Final configuration = model_unmasked_v5:blend0.3.")
arms_sheet("S10_Load", "step10_7", "Step 10.7: load-aware variants of the final model", "pen0.x = score penalty on busiest-10% specialists; cap = weekly capacity; tb = tie-break bands; v6_load / v5_posw = retrained variants.")
arms_sheet("S11_Base", "step11_1_tabular", "Step 11.1: graph model vs pure-tabular baselines", "lgbm_tabular = LightGBM on claim attributes only (no history, no network); logreg_tabular = logistic regression, same inputs.")

# =============================================================== Step 11 lens compare + groups ===========================================================
lc = J(f"{A}/step11_1_tabular/lens_compare.json")
rows = []
for yr, sp in (("2018", "validation"), ("2019", "test")):
    for k, v in lc[sp].items():
        if "|" not in k: continue
        var, lens = k.split("|"); rows.append({"year": yr, "setting": "30% quality weight" if var == "q30" else "model only", "lens": lens, "GNN": v["gnn"], "Table (LightGBM)": v["table"],
                                               "reported GNN − table": v["diff"], "ci_lo": v["ci"][0], "ci_hi": v["ci"][1]})
d = pd.DataFrame(rows)
ws = sheet("S11_GraphVsTable", "Step 11: graph model vs table on the four lenses (+ new to PCP)", "Seed-averaged (2 + 2 runs). PCP-cluster 95% intervals. Green: difference recomputed from the two columns and checked against the reported value.", f"{A}/step11_1_tabular/lens_compare.json")
fm = {c: "0.000" for c in ("GNN", "Table (LightGBM)", "reported GNN − table", "ci_lo", "ci_hi")}
cols = table(ws, d, 5, name="graph_vs_table", fmts=fm)
checkcol(ws, cols, len(d), lambda r: f"=D{r}-E{r}", "recomputed GNN − table", "I", "0.000"); checkcol(ws, cols, len(d), lambda r: f"=ABS(I{r}-F{r})<1E-9", "matches?", "J")
checkcol(ws, cols, len(d), lambda r: f'=IF(H{r}<0,"below 0",IF(G{r}>0,"above 0","crosses 0"))', "interval", "K")
for i, r in enumerate(rows): REG[("S11_GraphVsTable", f"{r['year']}|{r['setting']}|{r['lens']}")] = 6 + i

st = J(f"{A}/step11_2/strata_summary.json")
rows = []
for k, rr in st["splits"].items():
    for r in rr:
        if r["referrals"] == 0: continue
        row = {"split": k, "group": r["stratum"], "referrals": r["referrals"], "patients": r["patients"], "pcps": r["pcps"], "small (<100 patients)": r["small_cell"]}
        for a in ("GNN", "GNN, no message passing", "Table (LightGBM)", "PCP habit"):
            for l_ in ("fit", "quality", "load", "new", "travel"): row[f"{a}|{l_}"] = r[f"{a}|q30|{l_}"]
        row["Actual|load"] = r["Actual referral|base|load"]
        for l_ in ("load", "quality", "fit"):
            row[f"GNN−table|{l_}"] = r[f"GNN−table|q30|{l_}"]; row[f"GNN−table|{l_}|ci_lo"], row[f"GNN−table|{l_}|ci_hi"] = r[f"GNN−table|q30|{l_}|ci"]
        rows.append(row)
d = pd.DataFrame(rows)
ws = sheet("S11_Groups", "Step 11.2: graph model vs table by group (30% quality weight, 2018) — the 9 Oct forest plot", "Small cells (< 100 patients) were not judged. load = busiest-10% share of the top pick; quality = MIPS gain (points).", f"{A}/step11_2/strata_summary.json")
cols = table(ws, d, 5, name="groups11", fmts={c: ("0.00" if ("quality" in c and "|" in c) else None) for c in d.columns if "quality" in c})
lc_ = cols["GNN|load"][2]; tc_ = cols["Table (LightGBM)|load"][2]; dc_ = cols["GNN−table|load"][2]; nxt = L(len(d.columns) + 1)
checkcol(ws, cols, len(d), lambda r: f"=ABS(({lc_}{r}-{tc_}{r})-{dc_}{r})<1E-9", "load diff = GNN − table?", nxt)

# =============================================================== Step 12: group tests ===========================================================
g = J(f"{A}/step12_groups/groups_summary.json")
did = [r for r in g["did"] if r["variant"] == "q30"]
d = pd.DataFrame([{"split": r["split"], "group": r["group"], "referrals": r["n"], "pcps": r["pcps"], "graph − table, this group": r["d_group"], "graph − table, everyone else": r["d_rest"],
                   "difference (group − rest)": r["did_est"], "ci_lo": r["did_ci"][0], "ci_hi": r["did_ci"][1], "p (one group)": r["did_p"], "reported Holm p": r["did_p_holm"]} for r in did])
ws = sheet("S12_GroupTests", "Are the 85+ and early-tenure-PCP exceptions real? (2018, 30% quality weight)", "Load difference (graph − table, busiest-10% share of top pick) in the group minus in everyone else; PCP-cluster bootstrap. Holm adjusts for testing all 7 groups. Green: Holm recomputed by formula.", f"{A}/step12_groups/groups_summary.json")
fm = {c: "0.000" for c in d.columns if c not in ("split", "group", "referrals", "pcps")}
cols = table(ws, d, 5, name="group_tests", fmts=fm); nrow = len(d); pr = f"$J$6:$J${5 + nrow}"
checkcol(ws, cols, nrow, lambda r: f"=({nrow}-RANK(J{r},{pr},1)+1)*J{r}", "Holm step value (m−rank+1)·p", "L", "0.000")
checkcol(ws, cols, nrow, lambda r: f'=MIN(1,_xlfn.MAXIFS($L$6:$L${5 + nrow},{pr},"<="&J{r}))', "Holm p (recomputed)", "M", "0.000")
checkcol(ws, cols, nrow, lambda r: f"=ABS(M{r}-K{r})<0.002", "matches reported?", "N")
for i, r in enumerate(did): REG[("S12_GroupTests", r["group"])] = 6 + i
age = pd.DataFrame([{"patient age": k, "referrals": v["n"], "busy share of eligible set": v["available_busy"], "actual choices": v["actual_busy"], "Graph model": v["GNN|q30|load"],
                     "Table": v["Table|q30|load"], "GNN, no message passing": v["No-graph|q30|load"]} for k, v in g["age_profile"].items()])
r0 = 8 + nrow; ws.cell(row=r0 - 1, column=1, value="Top-pick load by patient age (the 'Load by patient age' chart)").font = FB
table(ws, age, r0, name="age_profile")
std = g["age_specialty_standardisation"]["q30"]; r1 = r0 + len(age) + 3
ws.cell(row=r1 - 1, column=1, value="Specialty-mix check: under-85 referrals reweighted to the 85+ specialty mix").font = FB
table(ws, pd.DataFrame([{"case": k, "Graph model": v.get("GNN"), "Table": v.get("Table"), "graph − table": v.get("d")} for k, v in std.items() if isinstance(v, dict)]), r1, name="age_std")
dose = pd.DataFrame([{"PCP history fifth": r["quintile"], "earlier referrals from": r["range"][0], "to": r["range"][1], "referrals": r["n"], "pcps": r["pcps"], "early-tenure referrals": r["early_tenure_pcp_referrals"],
                      "graph − table": r["d|q30"]["est"], "ci_lo": r["d|q30"]["ci"][0], "ci_hi": r["d|q30"]["ci"][1]} for r in g["pcp_history_dose"]["pcp_prior_refs"]])
r2 = r1 + 8; ws.cell(row=r2 - 1, column=1, value="By the PCP's referral history (fifths)").font = FB; table(ws, dose, r2, name="pcp_dose", fmts={"graph − table": "0.000", "ci_lo": "0.000", "ci_hi": "0.000"})

# =============================================================== Step 12: top-K ===========================================================
tk = J(f"{A}/step12_topk/topk_summary.json"); rows = []
for K in (5, 10, 15):
    T = tk[f"K{K}"]
    for m in ("fit", "load_top1", "load_slots", "new", "quality", "travel"):
        rows.append({"K": K, "lens": m, "GNN": T["arms"]["GNN|q30"][m], "Table": T["arms"]["Table (LightGBM)|q30"][m], "PCP habit": T["arms"]["PCP habit|q30"][m],
                     "reported GNN − table": T["gnn_minus_table"][f"{m}|q30"]["est"], "ci_lo": T["gnn_minus_table"][f"{m}|q30"]["ci"][0], "ci_hi": T["gnn_minus_table"][f"{m}|q30"]["ci"][1],
                     "lists filled beyond 25 mi": T["filled_beyond_25mi"]})
d = pd.DataFrame(rows)
ws = sheet("S12_TopK", "Graph model vs table with 5, 10 and 15 names (2018, 30% quality weight)", "load_top1 = busiest share of the first pick (independent of K); load_slots = share of all K list slots on the busiest 10%; quality = MIPS gain over the list; travel = mean miles.", f"{A}/step12_topk/topk_summary.json")
cols = table(ws, d, 5, name="topk", fmts={c: "0.000" for c in ("GNN", "Table", "PCP habit", "reported GNN − table", "ci_lo", "ci_hi")})
checkcol(ws, cols, len(d), lambda r: f"=ABS((C{r}-D{r})-F{r})<1E-9", "diff = GNN − table?", "K")
for i, r in enumerate(rows): REG[("S12_TopK", f"{r['K']}|{r['lens']}")] = 6 + i

# =============================================================== Step 12a mechanism ===========================================================
f12, s12, dc = J(f"{A}/step12a/final_summary.json"), J(f"{A}/step12a/secondary_summary.json"), J(f"{A}/step12a_decisions.json")
ws = sheet("S12a_Mechanism", "Why does the graph model ease congestion? (pre-registered STEP12a + post hoc amendment A1)", "Load = busiest-10% share of top pick at the 30% quality weight; effects in load points at matched fit (random-mixture dilution).", [f"{A}/step12a/final_summary.json", f"{A}/step12a/secondary_summary.json", f"{A}/step12a_decisions.json"])
gate = pd.DataFrame([{"quantity": "d_G: table − neural model with NO PCP history, at matched fit", "estimate": dc["d_G"]["est"], "ci_lo": dc["d_G"]["ci"][0], "ci_hi": dc["d_G"]["ci"][1], "seed pairings > 0": sum(x > 0 for x in dc["d_G"]["seed_pairings"]), "of": len(dc["d_G"]["seed_pairings"])},
                     {"quantity": "d_NG: table − GNN without message passing, at matched fit", "estimate": dc["d_NG"]["est"], "ci_lo": dc["d_NG"]["ci"][0], "ci_hi": dc["d_NG"]["ci"][1], "seed pairings > 0": None, "of": None},
                     {"quantity": "R (gap to explain, frozen)", "estimate": f12["R"], "ci_lo": None, "ci_hi": None, "seed pairings > 0": None, "of": None},
                     {"quantity": "pre-set neural run-to-run noise band", "estimate": f12["noise_band_neural"], "ci_lo": None, "ci_hi": None, "seed pairings > 0": None, "of": None}])
ws.cell(row=5, column=1, value="Gate and thresholds").font = FB; cols_g = table(ws, gate, 6, name="mech_gate", fmts={"estimate": "0.000", "ci_lo": "0.000", "ci_hi": "0.000"})
REG[("S12a", "dG")] = "'S12a_Mechanism'!$B$7"; REG[("S12a", "dG_lo")] = "'S12a_Mechanism'!$C$7"; REG[("S12a", "R")] = "'S12a_Mechanism'!$B$9"; REG[("S12a", "noise")] = "'S12a_Mechanism'!$B$10"
H, HS = f12["hypotheses"], s12["hypotheses"]; SEC = {"H1": "H1", "H2": "H2", "H3": "H3′", "H5": "H5"}
hyp = pd.DataFrame([{"test": h, "contrast": H[h]["contrast"], "effect": H[h]["effect"], "ci_lo": H[h]["ci"][0], "ci_hi": H[h]["ci"][1], "Holm p": H[h].get("p_holm"),
                     "all seed pairings > 0": H[h]["criteria"].get("all_seed_pairings"), "verdict (pre-registered)": H[h]["verdict"],
                     "post hoc effect": HS[SEC[h]]["effect"] if h in SEC else None, "post hoc verdict": HS[SEC[h]].get("verdict_secondary") if h in SEC else None} for h in ("H1", "H2", "H3", "H4", "H5")])
ws.cell(row=12, column=1, value="Pre-registered tests").font = FB; cols_h = table(ws, hyp, 13, name="mech_tests", fmts={c: "0.000" for c in ("effect", "ci_lo", "ci_hi", "Holm p", "post hoc effect")})
checkcol(ws, cols_h, len(hyp), lambda r: f"=AND(D{r}>0,C{r}>=$B$9/2)", "test criterion: effect ≥ R/2 and CI > 0 (outcome, not a check)", "L"); checkcol(ws, cols_h, len(hyp), lambda r: f"=C{r}>$B$10", "test criterion: beyond noise band (outcome, not a check)", "M")
REG[("S12a", "H1")] = "'S12a_Mechanism'!$C$14"
D_ = f12["descriptive"]
cells = pd.DataFrame([{"model": k, **{m: v["lenses_q30"][m] for m in ("fit", "load", "mips_gain", "new_to_pcp", "travel_mean_mi")}, "seeds": v["lenses_q30"].get("n_seeds"),
                       "top-pick volume percentile": v.get("volpct1_ft"), "load change with volume neutralised (matched)": (v.get("B1_no_volume") or {}).get("load_change_matched")}
                      for k, v in D_.items() if isinstance(v, dict) and "lenses_q30" in v])
r0 = 15 + len(hyp); ws.cell(row=r0, column=1, value="Model cells (seed-averaged), 30% quality weight").font = FB
table(ws, cells, r0 + 1, name="mech_cells", fmts={"mips_gain": "0.00", "travel_mean_mi": "0.00", "load change with volume neutralised (matched)": "0.000"})
sc = pd.DataFrame([{"cell": k, **{m: v[m] for m in ("family", "fit", "load", "mips_gain", "new_to_pcp", "seeds")}} for k, v in s12["cells"].items()])
r1 = r0 + len(cells) + 4; ws.cell(row=r1, column=1, value="Amendment A1 cells (incl. 6-seed neural model without history, volume-blind 'strict' cells)").font = FB; table(ws, sc, r1 + 1, name="mech_a1_cells")
cg = pd.DataFrame(f12["capacity_grid"]); r2 = r1 + len(sc) + 4; ws.cell(row=r2, column=1, value="LightGBM capacity grid (leaves × rounds)").font = FB; table(ws, cg, r2 + 1, name="mech_capacity")

# =============================================================== Step 12b hybrids ===========================================================
jb = J(f"{A}/step12b/judge_summary.json"); LK = ["r5", "busy1", "busy5", "new_pcp5", "new_pcp1", "gain5", "good_new5", "new_ins5", "median_travel_mi"]
NM = {"hyb_blend0.25": "Score blend, 25% table", "hyb_blend0.5": "Score blend, 50/50", "hyb_blend0.75": "Score blend, 75% table", "hyb_gate": "Gate by PCP / patient profile", "hyb_rrf": "Rank fusion (RRF)",
      "hyb_split1": "GNN top pick, table for the rest", "hyb_split2": "GNN top 2, table for the rest", "s12_lgbm_flat_lambdarank": "Flattened graph → LightGBM",
      "s12b_emb_lr": "GNN encoding + variables → logistic regression", "s12b_vars_lr": "Same variables, no GNN encoding"}
rows = [{"arm": k, **{m: v[m] for m in LK if m in v}} for k, v in jb["references"].items()]
for k, a in jb["arms"].items():
    row = {"arm": NM.get(k, k), **{m: a["lenses"].get(m) for m in LK}, "passes pre-stated rule": a["combines_strengths"]}
    for m in ("r5", "busy1", "busy5", "new_pcp5", "gain5"):
        c = a["ci_delta_vs_GNN"].get(m); row[f"Δ{m} vs GNN ci_lo"], row[f"Δ{m} vs GNN ci_hi"] = (c or [None, None])
    rows.append(row)
d = pd.DataFrame(rows)
ws = sheet("S12b_Hybrids", "Simple hybrids of the graph model and the table (2018, 30% quality weight, 2 seed pairings)", "Pre-stated rule: GNN's load (top pick and top 5) and quality, plus ≥ 5 points more new-to-PCP connections with CI > 0. Green: Δ vs GNN recomputed from the lens columns.", f"{A}/step12b/judge_summary.json")
cols = table(ws, d, 5, name="hybrids", fmts={c: "0.000" for c in d.columns if c.startswith("Δ")})
nx = len(d.columns)
for j, m in enumerate(("busy1", "new_pcp5")):
    cl = cols[m][2]; checkcol(ws, cols, len(d), lambda r, cl=cl: f"={cl}{r}-{cl}$6", f"Δ{m} vs GNN (formula)", L(nx + 1 + j), "0.000")
REG[("S12b", "n_arms")] = len(jb["arms"]); REG[("S12b", "pass_range")] = f"'S12b_Hybrids'!${cols['passes pre-stated rule'][2]}$8:${cols['passes pre-stated rule'][2]}${5 + len(d)}"

# =============================================================== Step 12c complementarity ===========================================================
cp = J(f"{A}/step12c/complementarity.json"); rows = []
for p, pp in enumerate(cp["per_pair"]):
    base = {"pairing": p + 1, "GNN run": cp["pairs"][p][0], "table run": cp["pairs"][p][1], "same top pick": pp["A_top1_agreement"], "top-5 overlap": pp["A_top5_overlap_mean"], "hindsight list-choice fit (upper bound)": pp["E2_hindsight_fit_oracle"]}
    for side, k in (("table only", "B_table_only_top5"), ("graph only", "B_gnn_only_top5")):
        for m in ("busy", "new", "ins", "cat", "mips", "dist"): base[f"{side}: {m}"] = pp[k][m]
    rows.append(base)
d = pd.DataFrame(rows)
ws = sheet("S12c_Complementarity", "Do the two models make different picks? (2018, 30% quality weight)", "'table only' / 'graph only' = specialists in one model's top 5 but not the other's, same referrals. busy = busiest 10%; ins = has seen the patient's insurer; cat = has treated the disease category; mips = MIPS score; dist = miles.", f"{A}/step12c/complementarity.json")
table(ws, d, 5, name="complementarity", fmts={c: "0.0" for c in d.columns if c.endswith("mips") or c.endswith("dist")})

# frontier points (12c G/T/B/S, 12d D/O/K)
FAM = {"G": "GNN", "T": "Table", "B": "Equal blend", "S": "Two-list selection (12c)", "D": "Dual encoder", "O": "Table-offset GNN", "K": "Stacked LightGBM"}
pts = []
for k in range(3):
    c12 = pd.read_csv(f"{A}/step12c/frontier_points_pair{k}.csv"); d12 = pd.read_csv(f"{A}/step12d/frontier_points_pair{k}.csv")
    d12 = d12[d12.family != "B"].assign(family=lambda x: x.family.map({"D": "D", "O": "O", "S": "K"}))
    for fr, stp in ((c12, "12c"), (d12, "12d")):
        for r in fr.itertuples(): pts.append({"step": stp, "pairing": k + 1, "family": r.family, "family name": FAM[r.family], "a (new bonus)": r.a_new, "b (busy penalty)": r.b_busy, "fit": r.fit,
                                              "load (top pick)": r.busy1, "load (top 5)": r.busy5, "new to PCP (top 5)": r.new5, "MIPS gain": r.gain5, "good new (top 5)": r.good_new5})
pdf = pd.DataFrame(pts)
ws = sheet("S12cd_FrontierPoints", "Steering grid: every (a, b) setting for every family and seed pairing (2018, 30% quality weight)", "a, b ∈ {0, .05, .1, .15, .2, .3, .4, .6}; a = b = 0 is the unsteered model. Used by the frontier sheets' MINIFS checks.", [f"{A}/step12c/frontier_points_pair{{0,1,2}}.csv", f"{A}/step12d/frontier_points_pair{{0,1,2}}.csv"])
cols = table(ws, pdf, 5, name="frontier_points", fmts={"MIPS gain": "0.00", "a (new bonus)": "0.00", "b (busy penalty)": "0.00"})
N = 5 + len(pdf); RNG = {k: f"'S12cd_FrontierPoints'!${cols[k][2]}$6:${cols[k][2]}${N}" for k in ("pairing", "family", "fit", "load (top pick)", "new to PCP (top 5)", "a (new bonus)", "b (busy penalty)")}

fr = J(f"{A}/step12c/frontiers.json"); fd = J(f"{A}/step12d/frontier.json")
ws = sheet("S12c_Frontier", "Lowest top-pick load each family can reach: fit ≥ 45% and new-to-PCP ≥ target (2018)", "Reported = value in the frontier JSON (GNN and table from the 12c file, stored to 4 decimals, so matched within 0.0001; the others exact); green = MINIFS over the 'S12cd_FrontierPoints' grid; blank = target unreachable at that fit.", [f"{A}/step12c/frontiers.json", f"{A}/step12d/frontier.json"])
rows = []; fam_order = ["G", "T", "B", "D", "O", "K"]
for k in range(3):
    for t in ("0.25", "0.3", "0.35", "0.4"):
        rep12c = fr["per_pair"][k]["floors"]["0.45"]["lowest_busy1_at_new5_target"][t]
        rep12d = fd["per_pair"][k]["floors"]["0.45"]["lowest_busy1_at_new5"][t]
        row = {"pairing": k + 1, "new-to-PCP target": float(t)}
        for f_ in fam_order:
            row[f"reported: {FAM[f_]}"] = (rep12c.get(f_) if f_ in "GT" else rep12d.get({"K": "S"}.get(f_, f_)))
        rows.append(row)
d = pd.DataFrame(rows); cols = table(ws, d, 5, name="frontier_lowest", fmts={c: "0.0%" for c in d.columns if c.startswith("reported")} | {"new-to-PCP target": "0%"})
for j, f_ in enumerate(fam_order):
    cl = L(len(d.columns) + 1 + j)
    checkcol(ws, cols, len(d), lambda r, f_=f_: f'=IFERROR(1/(1/_xlfn.MINIFS({RNG["load (top pick)"]},{RNG["pairing"]},A{r},{RNG["family"]},"{f_}",{RNG["fit"]},">=0.45",{RNG["new to PCP (top 5)"]},">="&B{r})),"")', f"MINIFS: {FAM[f_]}", cl, "0.0%")
REG[("S12c_Frontier", "first_row")] = 6
checkcol(ws, cols, len(d), lambda r: f'=AND(IF(C{r}="",I{r}="",ABS(I{r}-C{r})<1E-4),IF(D{r}="",J{r}="",ABS(J{r}-D{r})<1E-4),IF(E{r}="",K{r}="",ABS(K{r}-E{r})<1E-9),IF(F{r}="",L{r}="",ABS(L{r}-F{r})<1E-9),IF(G{r}="",M{r}="",ABS(M{r}-G{r})<1E-9),IF(H{r}="",N{r}="",ABS(N{r}-H{r})<1E-9))', "all families match reported?", L(len(d.columns) + 7))

# =============================================================== Steered blend 2018 / 2019 ===========================================================
pv, pt, pre = J(f"{A}/step12c/point_lenses_validation.json"), J(f"{A}/step12c/point_lenses_test2019.json"), J(f"{A}/step12c/preregistration_2019.json")
LM = ["fit", "busy1", "busy5", "new1", "new5", "gain5", "good_new5", "new_insurer_rate", "treated_category5", "dist_mean5", "dist_median_top1"]
rows = []
for yr, P in (("2018", pv), ("2019", pt)):
    for k, pp in enumerate(P["per_pair"]):
        for arm in ("GNN", "Table", "Steered blend"): rows.append({"year": yr, "pairing": k + 1, "arm": arm, **{m: pp[arm].get(m) for m in LM}})
d = pd.DataFrame(rows)
ws = sheet("S12c_SteeredBlend", "Steered equal blend vs graph model and table: 2018 (setting chosen) and 2019 (held out)", f"Setting: {json.dumps(pre['setting'], ensure_ascii=False)}. Pre-registered {pre['recorded']}. Rule: {pre['selection_rule_on_2018']}. Holding up: {pre['what_would_count_as_holding_up']}", [f"{A}/step12c/point_lenses_validation.json", f"{A}/step12c/point_lenses_test2019.json", f"{A}/step12c/preregistration_2019.json"])
cols = table(ws, d, 5, name="steered_blend", fmts={"gain5": "0.00", "dist_mean5": "0.00", "dist_median_top1": "0.00"})
REG[("Blend", "rows")] = {(r["year"], r["pairing"], r["arm"]): 6 + i for i, r in enumerate(rows)}
crit = []
for yr, P in (("2018", pv), ("2019", pt)):
    for k, pp in enumerate(P["per_pair"]):
        c = pp["delta_blend_minus_GNN_ci"]; crit.append({"year": yr, "pairing": k + 1, **{f"{m} ci_lo": c[m][0] for m in ("busy1", "busy5", "new5", "gain5", "fit")}, **{f"{m} ci_hi": c[m][1] for m in ("busy1", "busy5", "new5", "gain5", "fit")}})
dcr = pd.DataFrame(crit); r0 = 8 + len(d); ws.cell(row=r0 - 1, column=1, value="Blend − graph model, 95% intervals, and the pre-stated 2019 criteria (green formulas)").font = FB
colc = table(ws, dcr, r0, name="blend_criteria", fmts={c: "0.000" for c in dcr.columns if "ci" in c})
def fitcell(yr, k, arm): return f"$D${REG[('Blend', 'rows')][(yr, k, arm)]}"
base = len(dcr.columns)
checkcol(ws, colc, len(dcr), lambda r: f"=AND(H{r}<0,I{r}<0)", "load lower (top pick & top 5)?", L(base + 1))
checkcol(ws, colc, len(dcr), lambda r: f"=E{r}>0", "more new connections?", L(base + 2))
checkcol(ws, colc, len(dcr), lambda r: f"=F{r}>=-0.5", "MIPS not lower by > 0.5?", L(base + 3))
for i, cr in enumerate(crit):
    rr = r0 + 1 + i; yr, k = cr["year"], cr["pairing"]
    c_ = ws[f"{L(base + 4)}{rr}"]; c_.value = f"=INDEX($D:$D,{REG[('Blend','rows')][(yr,k,'Steered blend')]})>=MIN(INDEX($D:$D,{REG[('Blend','rows')][(yr,k,'GNN')]}),INDEX($D:$D,{REG[('Blend','rows')][(yr,k,'Table')]}))-0.01"; c_.fill = CHK; c_.font = F
    c2 = ws[f"{L(base + 5)}{rr}"]; c2.value = f"=AND({L(base + 1)}{rr},{L(base + 2)}{rr},{L(base + 3)}{rr},{L(base + 4)}{rr})"; c2.fill = CHK; c2.font = F
ws[f"{L(base + 4)}{r0}"] = "fit within 1 pt of weaker model?"; ws[f"{L(base + 5)}{r0}"] = "all criteria hold?"
for c in (f"{L(base + 4)}{r0}", f"{L(base + 5)}{r0}"): ws[c].font = FH; ws[c].fill = PatternFill("solid", fgColor="548235")
REG[("Blend", "all2019")] = f"'S12c_SteeredBlend'!${L(base + 5)}${r0 + 4}:${L(base + 5)}${r0 + 6}"

# =============================================================== Step 12d heavy ===========================================================
hd = J(f"{A}/step12d/heavy_descriptives.json"); rows = []
for k in range(3):
    for f_, v in fd["per_pair"][k]["unsteered"].items(): rows.append({"pairing": k + 1, "family": {"B": "Equal blend", "D": "Dual encoder", "O": "Table-offset GNN", "S": "Stacked LightGBM"}[f_], **v})
ws = sheet("S12d_Heavy", "Hybrids trained together: unsteered lenses and what each one learned (2018, 30% quality weight)", "Offset β = learned weight on the table's score (initialised at 1). Stacking gain share = share of LightGBM split gain on the GNN-score column.", [f"{A}/step12d/frontier.json", f"{A}/step12d/heavy_descriptives.json"])
table(ws, pd.DataFrame(rows), 5, name="heavy_unsteered", fmts={"gain5": "0.00", "dist5": "0.00"})
r0 = 8 + len(rows)
ws.cell(row=r0 - 1, column=1, value="Learned offset β and stacking gain share").font = FB
table(ws, pd.DataFrame([{"run": k, "offset β": v["beta"], "best epoch": v["best_epoch"]} for k, v in hd["offset_beta"].items()]), r0, name="heavy_beta", fmts={"offset β": "0.000"})
table(ws, pd.DataFrame([{"run": k, "GNN-score gain share": v["share"], "rank": v["rank"], "of features": v["n_features"]} for k, v in hd["stack_gnn_gain_share"].items()]), r0 + 6, name="heavy_gain", fmts={"GNN-score gain share": "0.0%"})

# =============================================================== Blend checks ===========================================================
bc = J(f"{A}/step12c/blend_checks.json"); rows = []
for s_, rr in bc["groups"].items():
    for r in rr:
        row = {"split": s_, "group": r["group"], "referrals": r["referrals"], "patients": r["patients"], "pcps": r["pcps"], "small (not judged)": r["small"]}
        for m in ("busy1", "new5", "gain5", "fit"):
            row[f"GNN|{m}"], row[f"Blend|{m}"], row[f"blend−GNN|{m}"] = r[f"GNN|{m}"], r[f"Blend|{m}"], r[f"blend−GNN|{m}"]
            row[f"blend−GNN|{m}|ci_lo"], row[f"blend−GNN|{m}|ci_hi"] = r[f"blend−GNN|{m}|ci"]
        rows.append(row)
d = pd.DataFrame(rows)
ws = sheet("S12c_BlendChecks", "Steered blend vs graph model by group, and where the freed top picks go (2018, 3 pairings averaged)", "Small = < 100 patients or < 5 PCPs. busy1 = top-pick load; new5 = new to PCP (top 5); gain5 = MIPS gain; fit = chosen in top 5.", f"{A}/step12c/blend_checks.json")
cols = table(ws, d, 5, name="blend_groups", fmts={c: ("0.00" if "gain5" in c else "0.000") for c in d.columns if "|" in c})
hi = cols["blend−GNN|busy1|ci_hi"][2]; sm = cols["small (not judged)"][2]; n_ = len(d)
REG[("BlendChecks", "n_judged")] = f'COUNTIF(\'S12c_BlendChecks\'!${sm}$6:${sm}${5 + n_},FALSE)'
REG[("BlendChecks", "n_lower")] = f'COUNTIFS(\'S12c_BlendChecks\'!${sm}$6:${sm}${5 + n_},FALSE,\'S12c_BlendChecks\'!${hi}$6:${hi}${5 + n_},"<0")'
e50 = [i for i, r in enumerate(rows) if r["group"] == "50+"][0]; REG[("BlendChecks", "fit50")] = f"'S12c_BlendChecks'!${cols['blend−GNN|fit'][2]}${6 + e50}"
REG[("BlendChecks", "fit50_hi")] = f"'S12c_BlendChecks'!${cols['blend−GNN|fit|ci_hi'][2]}${6 + e50}"
W = bc["where_the_load_goes"]; r0 = 8 + n_
wl = pd.DataFrame([{"top pick is…": lab, "Graph model": W["GNN"][k], "Table": W["Table"][k], "Steered blend": W["Blend"][k], "actual choices": W["Practice (chosen)"].get(k), "available (eligible sets)": W["Available (eligible sets)"].get(k.rstrip("1"))}
                   for lab, k in (("busiest 10%", "busy1"), ("early-career (≤ 6 Medicare years)", "early1"), ("low-experience (< 10 referrals)", "lowexp1"), ("non-metro", "nonmetro1"))])
ws.cell(row=r0 - 1, column=1, value="Where the load goes (top pick)").font = FB; table(ws, wl, r0, name="where_load")
tiers = pd.DataFrame([{"volume tier within the eligible set": t, "Graph model": W["GNN"]["top_pick_volume_tier"][t], "Table": W["Table"]["top_pick_volume_tier"][t], "Steered blend": W["Blend"]["top_pick_volume_tier"][t]} for t in W["GNN"]["top_pick_volume_tier"]])
cols_t = table(ws, tiers, r0 + 7, name="volume_tiers")
checkcol(ws, cols_t, 1, lambda r: f"=ROUND(SUM(B{r0 + 8}:B{r0 + 11}),6)=1", "tiers sum to 100%?", "E")

# =============================================================== Step 12e (top-K hybrids), if present ===========================================================
E = sorted(glob.glob(f"{A}/step12e/K*/step12e_K*_*.json"))
if E:
    rows = []
    for p in E:
        e = J(p)
        for k, pp in enumerate(e["per_pair"]):
            for arm, v in pp["fixed"].items():
                ci_ = pp["ci_vs_GNN"].get(arm, {})
                rows.append({"K": e["k"], "year": "2019" if e["split"] == "test" else "2018", "pairing": k + 1, "arm": arm, **v,
                             "Δload vs GNN ci_lo": (ci_.get("busy1") or [None, None])[0], "Δload vs GNN ci_hi": (ci_.get("busy1") or [None, None])[1],
                             "Δfit vs GNN ci_lo": (ci_.get("fit") or [None, None])[0], "Δfit vs GNN ci_hi": (ci_.get("fit") or [None, None])[1]})
    d = pd.DataFrame(rows).rename(columns={"busyK": "load (list slots)", "newK": "new to PCP (list)", "gainK": "MIPS gain (list)", "good_newK": "good new (list)", "distK": "mean travel (list)"})
    ws = sheet("S12e_TopK_Hybrids", "Steered blend and hybrids with 10 and 15 names (fixed 2018 setting, not re-tuned)", "Plan: STEP12E_TOPK_PLAN.md (frozen 2026-10-19). Stacked LightGBM excluded (its GNN-score lookup covers top-5 sets only). 2019 rows: GNN, table and blends only.", E)
    cols_e = table(ws, d, 5, name="topk_hybrids", fmts={"MIPS gain (list)": "0.00", "mean travel (list)": "0.00"} | {c: "0.000" for c in d.columns if "ci_" in c})
    REG[("S12e", "rng")] = {k: f"'S12e_TopK_Hybrids'!${cols_e[k][2]}$6:${cols_e[k][2]}${5 + len(d)}" for k in ("K", "year", "arm", "Δload vs GNN ci_hi", "Δfit vs GNN ci_lo", "Δfit vs GNN ci_hi", "pairing")}

# =============================================================== Step 13 / post hoc / Step 14 ===========================================================
V13, T13 = J(f"{A}/step13/eval/validation_step13_verdicts.json"), J(f"{A}/step13/eval/test_step13_verdicts.json")
rows = []
for yr, X in (("2018", V13), ("2019", T13)):
    for u, x in X["units"].items():
        for c_ in x["components"]:
            rows.append({"year": yr, "claim": PLAIN_CLAIM[u], "component": plain_component(c_["name"]), "point": c_["point"], "ci_lo": c_["ci95"][0], "ci_hi": c_["ci95"][1], "p_one_sided": c_["p_one_sided"],
                         **{f"seed {k}": v for k, v in c_["per_seed"].items()}, "undefined_share": c_["undefined_share"]})
d13 = pd.DataFrame(rows)
ws = sheet("S13_GraphValue", "What the graph adds: pre-registered claims and their comparisons (2018 decides; 2019 once)",
           "Pre-registered plan (frozen 2026-10-21; STEP13_GRAPH_VALUE_PLAN.md) with three changes made before any outcome was seen (STEP13_AMENDMENTS.md). Points are differences of proportions "
           "(0.026 = 2.6 points), graph side minus comparator; two-level paired bootstrap (seeds and PCPs), 2,000 draws. A claim's p = maximum over its comparisons; Holm across the 5 claims (2018).",
           [f"{A}/step13/eval/validation_step13_verdicts.json", f"{A}/step13/eval/test_step13_verdicts.json"])
cols13 = table(ws, d13, 5, name="s13_components", fmts={c: "0.000" for c in d13.columns if c not in ("year", "claim", "component")})
R13 = {k: f"'S13_GraphValue'!${cols13[k][2]}$6:${cols13[k][2]}${5 + len(d13)}" for k in ("year", "claim", "component", "point", "ci_lo", "ci_hi", "p_one_sided")}
r0u = 8 + len(d13)
du = pd.DataFrame([{"claim (2018)": PLAIN_CLAIM[u], "verdict": x["verdict"], "p_unit (file)": x["p_unit"], "p_holm (file)": x["p_holm"], "slice referrals": x.get("slice_n")} for u, x in V13["units"].items()])
colsu = table(ws, du, r0u, name="s13_units", fmts={"p_unit (file)": "0.000", "p_holm (file)": "0.000", "slice referrals": "0"})
cu = colsu["claim (2018)"][2]; pu = colsu["p_unit (file)"][2]
checkcol(ws, colsu, len(du), lambda r: f'=_xlfn.MAXIFS({R13["p_one_sided"]},{R13["year"]},"2018",{R13["claim"]},{cu}{r})', "p_unit recomputed (max of components)", L(len(du.columns) + 1), "0.000")
checkcol(ws, colsu, len(du), lambda r: f"=ABS({L(len(du.columns) + 1)}{r}-{pu}{r})<1e-9", "matches file?", L(len(du.columns) + 2))
REG[("S13", "verdicts")] = f"'S13_GraphValue'!${colsu['verdict'][2]}${r0u + 1}:${colsu['verdict'][2]}${r0u + len(du)}"
def comp_cell(yr, unit, name, col):
    i = d13.index[(d13.year == yr) & (d13.claim == PLAIN_CLAIM[unit]) & (d13.component == plain_component(name))][0]
    return f"'S13_GraphValue'!${cols13[col][2]}${6 + i}"
REG[("S13", "cc")] = comp_cell

PH = {yr: J(f"{A}/step13/posthoc_matched_load_{sp}.json") for yr, sp in (("2018", "validation"), ("2019", "test"))}
rows = []
for yr, X in PH.items():
    for r in X["pairs"]:
        L_, N_ = r["matched_top_pick_load_A_minus_B"], r["matched_newness_A_minus_B"]
        rows.append({"year": yr, "A": PLAIN_CELL[r["A"]], "B": PLAIN_CELL[r["B"]], "what": r["what"].replace(" (STEP11 headline)", " (earlier headline result)"), "fit A": r["fit_A"], "fit B": r["fit_B"], "Δload matched": L_["point"], "Δload ci_lo": L_["ci95"][0],
                     "Δload ci_hi": L_["ci95"][1], "Δnew matched": N_["point"], "Δnew ci_lo": N_["ci95"][0], "Δnew ci_hi": N_["ci95"][1]})
dph = pd.DataFrame(rows)
ws = sheet("S13_MatchedFit", "Post hoc: top-pick load and newness at matched fit (descriptive)",
           "Run after the 2019 pass on the graph-value score files; both years already seen, so descriptive, not a test. The more accurate model is mixed with the exact random scorer down to the other's fit "
           "(the matching used in the congestion study); models seed-averaged; PCP-cluster bootstrap, 2,000 draws. A − B; negative load = A spreads load more.",
           [f"{A}/step13/posthoc_matched_load_validation.json", f"{A}/step13/posthoc_matched_load_test.json"])
colsph = table(ws, dph, 5, name="s13_matched", fmts={c: "0.000" for c in dph.columns if c not in ("year", "A", "B", "what")})
REG[("PH", "cell")] = lambda yr, a, b, col: f"'S13_MatchedFit'!${colsph[col][2]}${6 + dph.index[(dph.year == yr) & (dph.A == PLAIN_CELL[a]) & (dph.B == PLAIN_CELL[b])][0]}"

C14, R14 = J(f"{A}/step14/choice_2018.json"), J(f"{A}/step14/results_2019.json")
rows = []
for f_, c_ in C14["choice"].items():
    base = {"family": PLAIN_FAM[f_], "proposed on 2018": f_ == C14["proposed"]}
    if c_.get("cannot_reach_floors"): rows.append({**base, "a_new": None, "b_busy": None, "year": "2018", "pairing": "cannot reach the floors"}); continue
    for yr, src in (("2018", c_["lenses_per_pairing"]), ("2019", R14["families"][f_]["lenses_per_pairing"])):
        for pk, lv in src.items():
            rows.append({**base, "a_new": c_["a_new"], "b_busy": c_["b_busy"], "year": yr, "pairing": f"s{pk}", **{k: lv[k] for k in ("fit", "busy1", "new5", "gain5")}})
d14 = pd.DataFrame(rows)
ws = sheet("S14_Blend", "Steered recommender with the graph-feature table: 2018 choice, 2019 check (per seed pairing)",
           "Pre-registered plan (frozen 2026-10-23; STEP14_GRAPH_FEATURE_BLEND_PLAN.md). Candidates: GNN, pure table, table + graph summaries, the earlier steered blend (GNN + table), "
           "GNN + table with graph summaries, and the three-way blend. Setting per candidate chosen on 2018: fit ≥ 0.45 and new ≥ 0.35 in all pairings, lowest top-pick load (busy1). "
           "Disclosed: motivated by the graph-value results, which had seen 2019.",
           [f"{A}/step14/choice_2018.json", f"{A}/step14/results_2019.json"])
cols14 = table(ws, d14, 5, name="s14_blend", fmts={"fit": "0.0%", "busy1": "0.0%", "new5": "0.0%", "gain5": "0.00"})
R14r = {k: f"'S14_Blend'!${cols14[k][2]}$6:${cols14[k][2]}${5 + len(d14)}" for k in ("family", "year", "pairing", "fit", "busy1", "new5")}
r0c = 8 + len(d14)
dcmp = pd.DataFrame([{"pairing": f"s{k}"} for k in (42, 1, 7)])
colsc = table(ws, dcmp, r0c, name="s14_check")
pc_ = colsc["pairing"][2]
def sumif(fam, col, r): return f'SUMIFS({R14r[col]},{R14r["family"]},"{PLAIN_FAM[fam]}",{R14r["year"]},"2019",{R14r["pairing"]},{pc_}{r})'
checkcol(ws, colsc, 3, lambda r: f"={sumif('B_GT', 'busy1', r)}-{sumif('B_GTF', 'busy1', r)}", "2019 load: earlier blend − three-way", "B", "0.000")
checkcol(ws, colsc, 3, lambda r: f"={sumif('B_GTF', 'fit', r)}-{sumif('B_GT', 'fit', r)}", "2019 fit: three-way − earlier blend", "C", "0.000")
checkcol(ws, colsc, 3, lambda r: f"={sumif('B_GTF', 'new5', r)}-{sumif('B_GT', 'new5', r)}", "2019 new: three-way − earlier blend", "D", "0.000")
checkcol(ws, colsc, 3, lambda r: f"=AND(B{r}>0,C{r}>=-0.01,D{r}>=-0.01)", "pre-stated test met?", "E")
REG[("S14", "test_rng")] = f"'S14_Blend'!$E${r0c + 1}:$E${r0c + 3}"
bt = R14["comparison_with_B_GT"]["proposed_minus_B_GT_paired_bootstrap"]
r0b = r0c + 6
dbt = pd.DataFrame([{"lens (three-way − earlier blend, 2019, paired PCP bootstrap)": k, "point": v["point"], "ci_lo": v["ci95"][0], "ci_hi": v["ci95"][1]} for k, v in bt.items()])
colsb = table(ws, dbt, r0b, name="s14_boot", fmts={"point": "0.000", "ci_lo": "0.000", "ci_hi": "0.000"})
REG[("S14", "load_ci_hi")] = f"'S14_Blend'!${colsb['ci_hi'][2]}${r0b + 1 + list(bt).index('busy1')}"
REG[("S14", "load_pt")] = f"'S14_Blend'!${colsb['point'][2]}${r0b + 1 + list(bt).index('busy1')}"

# =============================================================== Columns dictionary ===========================================================
COLS_DOC = [("r5 / fit", "Share of referrals with the chosen specialist in the top 5 (top K)."), ("r10", "Same for the top 10."), ("rr", "Mean reciprocal rank of the chosen specialist (sanity only)."),
            ("gain5", "Mean MIPS of the top 5 minus MIPS of the chosen specialist (points)."), ("gain5_cons", "Cautious bound: unscored/ambiguous MIPS counted against the list."), ("gain1", "MIPS of the top pick minus the chosen."),
            ("busy1", "Top pick is a busiest-10% specialist (share of referrals)."), ("busy5", "Share of the 5 slots on busiest-10% specialists."), ("new_pcp5 / new5", "Share of the 5 slots the PCP never referred to before."),
            ("new_pcp1 / new1", "Top pick is new to the PCP."), ("good_new5", "Slots that are new to the PCP, higher MIPS than the chosen, and not busiest."), ("new_ins5", "Slots new to the PCP whose specialist has seen the patient's insurer."),
            ("new_insurer_rate", "Among new-to-PCP slots, share whose specialist has seen the patient's insurer."), ("insurer5 / insurer_seen5", "Slots whose specialist has seen the patient's insurer before."),
            ("category5 / treated_category5", "Slots whose specialist has treated the disease category before."), ("dist5_median / dist_mean5", "Distance from the patient to the recommended specialists (median / mean, miles)."),
            ("within1_1", "Top pick's MIPS within 1 point of the chosen or better."), ("gini_within", "Gini of recommendations across specialists within specialty (concentration)."),
            ("_lo / _hi / ci_lo / ci_hi", "95% interval bounds."), ("q30 / blend0.3", "30% MIPS weight in the ranking."), ("resortK", "Model's top K re-sorted by MIPS."), ("pen0.x", "Score penalty on busiest-10% specialists (Step 10.7).")]
ws = sheet("Columns", "Column dictionary", "Metric names used across the sheets.", "STEP10_PROTOCOL.md §2.4; scripts/evaluate/step10_evaluate.py")
table(ws, pd.DataFrame(COLS_DOC, columns=["column", "meaning"]), 5, name="columns_doc")

# =============================================================== Claims ===========================================================
ws = sheet("Claims", "Headline claims in the decks and the cells that verify them", "Each row: the statement, where it appears, the supporting value(s) computed from the evidence sheets, and a TRUE/FALSE check.", "This workbook")
def S(sh, row, col): return f"'{sh}'!${col}${row}"
gt = lambda yr, lens: REG[("S11_GraphVsTable", f"{yr}|30% quality weight|{lens}")]
s10 = REG[("S10_Final_2019", "_cols")]; i_fin = REG[("S10_Final_2019", "model_unmasked_v5:blend0.3")]; i_pcp = REG[("S10_Final_2019", "pcp:base")]; i_pr = REG[("S10_Final_2019", "practice:base")]
s10v = REG[("S10_Final_2018", "_cols")]; i_base = REG[("S10_Final_2018", "model_unmasked_v5:base")]; i_prv = REG[("S10_Final_2018", "practice:base")]
s107 = REG[("S10_Load_2019", "_cols")]; i_pen = REG[("S10_Load_2019", "model_unmasked_v5:pen0.2")]; i_pr7 = REG[("S10_Load_2019", "practice:base")]
c = lambda sh, cols, col, i: cell_of(wb[sh], cols, col, i)
CL = [
    ("Final model with 30% quality weight raises MIPS on 2019 (points over the chosen specialist)", "2 Oct deck", f"={c('S10_Final_2019', s10, 'gain5', i_fin)}", f"={c('S10_Final_2019', s10, 'gain5_lo', i_fin)}>0"),
    ("…at fit close to the PCP's own habit (2019)", "2 Oct deck", f"={c('S10_Final_2019', s10, 'r5', i_fin)}-{c('S10_Final_2019', s10, 'r5', i_pcp)}", f"=ABS({c('S10_Final_2019', s10, 'r5', i_fin)}-{c('S10_Final_2019', s10, 'r5', i_pcp)})<0.02"),
    ("Imitation model concentrates top picks on the busiest 10% more than practice (2018, no quality weight)", "2 Oct deck", f"={c('S10_Final_2018', s10v, 'busy1', i_base)}", f"={c('S10_Final_2018', s10v, 'busy1', i_base)}>{c('S10_Final_2018', s10v, 'busy1', i_prv)}"),
    ("Load penalty 0.2 brings top-pick load below practice on 2019", "2 Oct deck", f"={c('S10_Load_2019', s107, 'busy1', i_pen)}", f"={c('S10_Load_2019', s107, 'busy1', i_pen)}<{c('S10_Load_2019', s107, 'busy1', i_pr7)}"),
    ("Graph model has lower top-pick load than the table, 2018 (GNN − table, CI below 0)", "9 Oct deck", f"={S('S11_GraphVsTable', gt('2018', 'load'), 'I')}", f"={S('S11_GraphVsTable', gt('2018', 'load'), 'H')}<0"),
    ("…and on 2019", "9 Oct deck", f"={S('S11_GraphVsTable', gt('2019', 'load'), 'I')}", f"={S('S11_GraphVsTable', gt('2019', 'load'), 'H')}<0"),
    ("Fit is level between graph model and table (2018 CI includes 0)", "9 Oct deck", f"={S('S11_GraphVsTable', gt('2018', 'fit'), 'I')}", f"=AND({S('S11_GraphVsTable', gt('2018', 'fit'), 'G')}<0,{S('S11_GraphVsTable', gt('2018', 'fit'), 'H')}>0)"),
    ("Table proposes more new-to-PCP connections (2018, CI below 0 for GNN − table)", "9 Oct deck", f"={S('S11_GraphVsTable', gt('2018', 'new'), 'I')}", f"={S('S11_GraphVsTable', gt('2018', 'new'), 'H')}<0"),
    ("85+ is not a real exception after adjusting for 7 groups (Holm p recomputed)", "16 Oct deck", f"={S('S12_GroupTests', REG[('S12_GroupTests', '85+')], 'M')}", f"={S('S12_GroupTests', REG[('S12_GroupTests', '85+')], 'M')}>0.05"),
    ("Early-tenure PCPs are not a real exception (Holm p recomputed)", "16 Oct deck", f"={S('S12_GroupTests', REG[('S12_GroupTests', 'early')], 'M')}", f"={S('S12_GroupTests', REG[('S12_GroupTests', 'early')], 'M')}>0.05"),
    ("Top-pick load gap does not change with list length (K = 5 vs 15)", "16 Oct deck", f"={S('S12_TopK', REG[('S12_TopK', '15|load_top1')], 'F')}-{S('S12_TopK', REG[('S12_TopK', '5|load_top1')], 'F')}", f"=ABS({S('S12_TopK', REG[('S12_TopK', '15|load_top1')], 'F')}-{S('S12_TopK', REG[('S12_TopK', '5|load_top1')], 'F')})<0.001"),
    ("List-slot load gap shrinks from top 5 to top 15", "16 Oct deck", f"={S('S12_TopK', REG[('S12_TopK', '15|load_slots')], 'F')}", f"=ABS({S('S12_TopK', REG[('S12_TopK', '15|load_slots')], 'F')})<ABS({S('S12_TopK', REG[('S12_TopK', '5|load_slots')], 'F')})"),
    ("A neural model with NO PCP-history inputs still has lower load than the table at matched fit", "16 Oct deck", f"={REG[('S12a', 'dG')]}", f"={REG[('S12a', 'dG_lo')]}>0"),
    ("Model-type effect (+3.9) is below the pre-set noise band → not shown under the pre-registered rule", "16 Oct deck", f"={REG[('S12a', 'H1')]}", f"={REG[('S12a', 'H1')]}<{REG[('S12a', 'noise')]}"),
    ("No simple hybrid passes the pre-stated 'combines the strengths' rule", "16 Oct deck", f"=COUNTIF({REG[('S12b', 'pass_range')]},TRUE)", f"=COUNTIF({REG[('S12b', 'pass_range')]},TRUE)=0"),
    ("Steered blend: lowest top-pick load at new ≥ 35%, fit ≥ 45%, pairing 1 (MINIFS over the grid; matches the frontier file)", "16 Oct deck", "='S12c_Frontier'!$K$8", "=AND(ISNUMBER('S12c_Frontier'!$K$8),'S12c_Frontier'!$O$8)"),
    ("The graph model alone cannot reach 35% new-to-PCP at fit ≥ 45% (any pairing)", "16 Oct deck", f'=COUNTIFS({RNG["family"]},"G",{RNG["fit"]},">=0.45",{RNG["new to PCP (top 5)"]},">=0.35")', f'=COUNTIFS({RNG["family"]},"G",{RNG["fit"]},">=0.45",{RNG["new to PCP (top 5)"]},">=0.35")=0'),
    ("Steered blend reaches lower load than every other family at new ≥ 35%, fit ≥ 45%, in all pairings", "16 Oct deck", "=MAX('S12c_Frontier'!$K$8,'S12c_Frontier'!$K$12,'S12c_Frontier'!$K$16)",
     "=AND(" + ",".join(f"'S12c_Frontier'!$K${r}<MIN('S12c_Frontier'!$I${r}:$J${r},'S12c_Frontier'!$L${r}:$N${r})" for r in (8, 12, 16)) + ")"),
    ("The pre-stated 2019 criteria hold in all three pairings", "16 Oct deck", f"=COUNTIF({REG[('Blend', 'all2019')]},TRUE)", f"=COUNTIF({REG[('Blend', 'all2019')]},TRUE)=3"),
    ("Steered blend lowers top-pick load vs the graph model in every judged group (CI below 0)", "16 Oct deck", f"={REG[('BlendChecks', 'n_lower')]}", f"={REG[('BlendChecks', 'n_lower')]}={REG[('BlendChecks', 'n_judged')]}"),
    ("Fit cost where the chosen specialist had 50+ earlier referrals (CI below 0)", "16 Oct deck", f"={REG[('BlendChecks', 'fit50')]}", f"={REG[('BlendChecks', 'fit50_hi')]}<0"),
    ("Table-offset GNN keeps the table at full weight (learned β ≈ 1)", "16 Oct deck", "=AVERAGE('S12d_Heavy'!$B$21:$B$23)", "=AND(COUNT('S12d_Heavy'!$B$21:$B$23)=3,ABS(AVERAGE('S12d_Heavy'!$B$21:$B$23)-1)<0.05)"),
]
if ("S12e", "rng") in REG:
    R_ = REG[("S12e", "rng")]
    for K in (10, 15):
        for yr in ("2018", "2019"):
            cnt = f'COUNTIFS({R_["K"]},{K},{R_["year"]},"{yr}",{R_["arm"]},"Steered blend")'
            low = f'COUNTIFS({R_["K"]},{K},{R_["year"]},"{yr}",{R_["arm"]},"Steered blend",{R_["Δload vs GNN ci_hi"]},"<0")'
            CL.append((f"Top {K}, {yr}: steered blend has lower top-pick load than the graph model in every pairing (CI below 0)", "16 Oct follow-up: longer lists", f"={low}", f"=AND({cnt}=3,{low}=3)"))
cc, ph = REG[("S13", "cc")], REG[("PH", "cell")]
CL += [
    ("No claim of the graph-value study is supported on 2018 (pre-registered, Holm over 5)", "Story deck (23 Oct)", f'=COUNTIF({REG[("S13", "verdicts")]},"supported")', f'=COUNTIF({REG[("S13", "verdicts")]},"supported")=0'),
    ("Message passing recovers less than half of what history columns give (the 'at least half' comparison is below 0, 2018)", "Story deck (23 Oct)", f"={cc('2018', 'H1a', 'joint half-gap', 'point')}", f"={cc('2018', 'H1a', 'joint half-gap', 'ci_hi')}<0"),
    ("Fixed 2-hop summaries in LightGBM beat the GNN without history, 2018 and 2019 (CI below 0)", "Story deck (23 Oct)", f"={cc('2018', 'H1b', 'GNN-nohist − flat_rf', 'point')}",
     f"=AND({cc('2018', 'H1b', 'GNN-nohist − flat_rf', 'ci_hi')}<0,{cc('2019', 'H1b', 'GNN-nohist − flat_rf', 'ci_hi')}<0)"),
    ("An MLP on the same summaries ties the GNN without history (2018 CI includes 0)", "Story deck (23 Oct)", f"={cc('2018', 'H1b', 'GNN-nohist − mlp_rf', 'point')}",
     f"=AND({cc('2018', 'H1b', 'GNN-nohist − mlp_rf', 'ci_lo')}<0,{cc('2018', 'H1b', 'GNN-nohist − mlp_rf', 'ci_hi')}>0)"),
    ("Post hoc: graph summaries cut a table's top-pick load at matched fit, 2018 and 2019 (CI below 0)", "Story deck (23 Oct)", f"={ph('2019', 'flat_rf', 'table', 'Δload matched')}",
     f"=AND({ph('2018', 'flat_rf', 'table', 'Δload ci_hi')}<0,{ph('2019', 'flat_rf', 'table', 'Δload ci_hi')}<0)"),
    ("Three-way blend meets the pre-stated 2019 test against the earlier steered blend in all 3 pairings (recomputed)", "Story deck (23 Oct)", f"=COUNTIF({REG[('S14', 'test_rng')]},TRUE)", f"=COUNTIF({REG[('S14', 'test_rng')]},TRUE)=3"),
    ("Three-way blend: lower 2019 top-pick load than the earlier steered blend (paired CI below 0)", "Story deck (23 Oct)", f"={REG[('S14', 'load_pt')]}", f"={REG[('S14', 'load_ci_hi')]}<0"),
]
d = pd.DataFrame([{"#": i + 1, "claim": a, "where": b} for i, (a, b, _, _) in enumerate(CL)])
cols = table(ws, d, 5, name="claims"); ws.column_dimensions["B"].width = 95
ws["D5"] = "value from evidence"; ws["E5"] = "check"
for cc in ("D5", "E5"): ws[cc].font = FH; ws[cc].fill = PatternFill("solid", fgColor="548235")
for i, (_, _, v, chk) in enumerate(CL):
    ws[f"D{6 + i}"] = v; ws[f"D{6 + i}"].fill = CHK; ws[f"D{6 + i}"].number_format = "0.000"; ws[f"E{6 + i}"] = chk; ws[f"E{6 + i}"].fill = CHK
ws.column_dimensions["D"].width = 18; ws.column_dimensions["E"].width = 10
ws[f"B{8 + len(CL)}"] = "All checks TRUE?"; ws[f"B{8 + len(CL)}"].font = FB; ws[f"E{8 + len(CL)}"] = f"=COUNTIF(E6:E{5 + len(CL)},TRUE)={len(CL)}"; ws[f"E{8 + len(CL)}"].fill = KEY

# sheet index on README, sheet order
order = ["README", "Claims", "Columns", "Cohort", "Decisions", "Runs"] + [s for s in wb.sheetnames if s.startswith("S1")]
wb._sheets = [wb[s] for s in order]
ws = wb["README"]; ws.cell(row=SHEET_INDEX_ROW + 1, column=1, value="Sheets").font = FB
for i, s in enumerate(order[1:]):
    ws.cell(row=SHEET_INDEX_ROW + 2 + i, column=1, value=s).font = F
    ws.cell(row=SHEET_INDEX_ROW + 2 + i, column=2, value=wb[s]["A1"].value).font = F
# =============================================================== plain-language pass ===========================================================
SHEET_NAMES = {"S10_Final_2018": "Quality weight 2018", "S10_Final_2019": "Quality weight 2019", "S10_Load_2018": "Load penalty 2018", "S10_Load_2019": "Load penalty 2019",
               "S11_Base_2018": "Baselines 2018", "S11_Base_2019": "Baselines 2019", "S11_GraphVsTable": "Graph vs table", "S11_Groups": "Graph vs table by group",
               "S12_GroupTests": "Group exceptions", "S12_TopK": "Longer lists", "S12a_Mechanism": "Why less concentrated", "S12b_Hybrids": "Simple hybrids",
               "S12c_Complementarity": "Complementarity", "S12cd_FrontierPoints": "Steering grid", "S12c_Frontier": "Steering frontier", "S12c_SteeredBlend": "Steered blend",
               "S12d_Heavy": "Heavier hybrids", "S12c_BlendChecks": "Steered blend checks", "S12e_TopK_Hybrids": "Steered blend longer lists",
               "S13_GraphValue": "What the graph adds", "S13_MatchedFit": "Matched fit (post hoc)", "S14_Blend": "Blend with graph summaries"}
TEXT_REPL = [("Step 10.3: final model (v5) vs heuristics, all re-rankings", "Graph model vs heuristics, all re-rankings (incl. the 30% quality weight)"),
             ("Step 10.7: load-aware variants of the final model", "Load-aware variants of the graph model"),
             ("Step 11.1: graph model vs pure-tabular baselines", "Graph model vs pure-tabular baselines"),
             ("Step 11: graph model vs table on the four lenses", "Graph model vs table on the four lenses"),
             ("Step 11.2: graph model vs table by group", "Graph model vs table by group"),
             ("(pre-registered STEP12a + post hoc amendment A1)", "(pre-registered tests + a disclosed post hoc addition)"),
             ("Amendment A1 cells (", "Post hoc addition ("), ("GNN and table from the 12c file", "GNN and table from the steered-blend frontier file"),
             ("Plan: STEP12E_TOPK_PLAN.md (frozen 2026-10-19).", "Pre-registered plan (frozen 2026-10-19; STEP12E_TOPK_PLAN.md)."),
             ("Source: STEP10_PROTOCOL.md §2.4;", "Source: the evaluation protocol (STEP10_PROTOCOL.md);"), ("(Step 10.7).", "(load-aware variants)."),
             ("Counts per exclusion step (STEP10_PROTOCOL §2.1).", "Counts per exclusion step (evaluation protocol)."),
             ("fixed STEP12c setting", "fixed steered-blend setting"), ("(no unit supported; H3 inconclusive)", "(no claim supported; insurer-compatibility claim inconclusive)")]
MECH = {"H1": "Model class (tree vs neural)", "H2": "Training objective", "H3": "Volume attributes", "H4": "Tree capacity", "H5": "Volume scaling"}
DEC_STEP = {"10.2": "Graph model on the lenses (first evaluation)", "10.3": "Final graph model with the quality weight", "10.4:reliance": "Reliance diagnostics",
            "10.4:cap10": "Distance cap 10 mi", "10.4:cap50": "Distance cap 50 mi", "10.4:capnone": "No distance cap", "10.4:ever_active": "Any earlier claim counts as active",
            "10.4:insurer_set": "Eligible set limited to the patient's insurer", "10.4:non_choice": "Facility-based specialties", "10.4:non_specific": "Non-specific specialties",
            "10.5": "Patient-group gaps", "10.5:capnone": "Patient-group gaps, no distance cap", "10.7": "Load-aware variants", "10.4:include_seen": "Already-seen specialists kept as candidates",
            "10.9": "Race concordance", "11.1-tabular": "Graph model vs pure-tabular baselines", "step12c_dump": "Steered blend (held-out check)",
            "step12e_topk_K10": "Steered blend with 10 names", "step12e_topk_K15": "Steered blend with 15 names", "13": "What the graph adds"}
def plain_amend(t):
    t = re.sub(r"STEP10_PROTOCOL A\d+(?: \+ A\d+)?", "evaluation-protocol amendment", t)
    t = t.replace("STEP11_PLAN G1 + B4", "graph-vs-table plan").replace("STEP12c (2026-10-14):", "steered-blend plan (2026-10-14):")
    t = t.replace("STEP12E_TOPK_PLAN.md (frozen 2026-10-19)", "longer-lists plan (frozen 2026-10-19)")
    return t.replace("STEP13_GRAPH_VALUE_PLAN.md (frozen 2026-10-21) + STEP13_AMENDMENTS.md A1, A2, A3", "graph-value plan (frozen 2026-10-21) + three pre-outcome changes")
for ws_ in wb.worksheets:
    for row in ws_.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str): continue
            if v.startswith("="):
                for old, new in SHEET_NAMES.items(): v = v.replace(f"'{old}'!", f"'{new}'!")
                cell.value = v; continue
            for old, new in TEXT_REPL: v = v.replace(old, new)
            for old, new in SHEET_NAMES.items():
                v = re.sub(rf"(?<![\w']){re.escape(old)}(?![\w])", new, v); v = v.replace(f"'{old}'", f"'{new}'")
            if ws_.title == "S12a_Mechanism" and v in MECH: v = MECH[v]
            if ws_.title == "Decisions":
                if cell.column == 1 and v in DEC_STEP: v = DEC_STEP[v]
                if cell.column == 4: v = {"model_s13_gnn_nohist:base": "GNN without history", "model_unmasked:blend0.3": "graph model, 30% quality weight",
                                          "model_unmasked_v5:blend0.3": "graph model, 30% quality weight", "model_unmasked_v5:pen0.2": "graph model, load penalty 0.2"}.get(v, v)
                v = plain_amend(v)
            cell.value = v
for old, new in SHEET_NAMES.items(): wb[old].title = new
os.makedirs(os.path.dirname(args.out), exist_ok=True); wb.save(args.out); print("saved", args.out, "sheets:", len(order))
