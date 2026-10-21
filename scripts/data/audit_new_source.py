"""Aggregate-only audit of a new source workbook, in the style of the cleaning notebook's
header audit. Prints: sheet names and sizes, header differences vs the known 119-column
source, and for location-like columns coverage and FORMAT PATTERNS only (never values).
Also reports patient-ID overlap COUNTS with the current prepared bundle.

Usage: audit_new_source.py <xlsx> [--sheet NAME] [--max-rows N]
"""
import os, sys, pathlib  # bootstrap: make referral_gnn importable from any scripts/ subfolder (or `pip install -e .`)
sys.path.insert(0, str(next(p for p in pathlib.Path(__file__).resolve().parents if (p / "referral_gnn" / "paths.py").is_file())))
from referral_gnn.paths import workspace, script_path, code_file  # noqa: E402,F401
import argparse
import os
import re
import sys
from collections import Counter

import openpyxl
import pandas as pd


root = str(workspace())

ap = argparse.ArgumentParser()
ap.add_argument("xlsx")
ap.add_argument("--sheet", default=None)
ap.add_argument("--max-rows", type=int, default=None)
args = ap.parse_args()

wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
print("sheets:", [(ws.title, ws.max_row, ws.max_column) for ws in wb.worksheets])
sheet = args.sheet or next((ws.title for ws in wb.worksheets if ws.max_row and ws.max_row > 1000), wb.worksheets[0].title)
ws = wb[sheet]
print("auditing sheet:", sheet)

rows = ws.iter_rows(values_only=True)
headers = [str(h).strip() if h is not None else "" for h in next(rows)]
print("columns:", len(headers))

# known headers from the original source (from the prepared bundle's manifest column lists where available)
known_path = os.path.join(root, "reviews", "known_headers_119.txt")
known = set(open(known_path).read().split("\n")) - {""} if os.path.exists(known_path) else None
if known:
    print("new vs known: added", sorted(set(headers) - known)); print("           missing", sorted(known - set(headers)))
else:
    print("(no known-header list at reviews/known_headers_119.txt; listing all headers)")
    print(headers)

loc_re = re.compile(r"(zip|zcta|postal|addr|address|city|county|state|lat|lon|geo|fips|ruca|dist|proxim)", re.I)
loc_cols = [i for i, h in enumerate(headers) if loc_re.search(h)]
print("\nlocation-like columns:", [headers[i] for i in loc_cols])

id_col = next((i for i, h in enumerate(headers) if h == "mso_membership_id"), None)
link_col = next((i for i, h in enumerate(headers) if h.upper().replace(" ", "") in ("DOC-DOC-LINK", "DOC_DOC_LINK", "DOCDOCLINK")), None)

def pattern(v):
    if v is None or (isinstance(v, str) and not v.strip()): return "missing"
    if isinstance(v, float) and v != v: return "missing"
    s = str(v).strip()
    if isinstance(v, (int, float)) and float(v).is_integer(): s = str(int(v))
    if re.fullmatch(r"\d{5}", s): return "5 digits"
    if re.fullmatch(r"\d{4}", s): return "4 digits (leading zero lost?)"
    if re.fullmatch(r"\d{3}", s): return "3 digits"
    if re.fullmatch(r"\d{9}", s): return "9 digits"
    if re.fullmatch(r"\d{5}-\d{4}", s): return "ZIP+4"
    if re.fullmatch(r"-?\d+\.\d+", s): return "decimal"
    if re.fullmatch(r"\d+", s): return f"{len(s)} digits"
    if re.fullmatch(r"[A-Za-z .'-]+", s): return "alphabetic"
    return "other"

n = 0; n_pos = 0
pat = {i: Counter() for i in loc_cols}; pat_pos = {i: Counter() for i in loc_cols}
ids = set(); ids_pos = set()
for r in rows:
    n += 1
    pos = link_col is not None and str(r[link_col]).strip() in ("1", "1.0", "True")
    if pos: n_pos += 1
    if id_col is not None and r[id_col] is not None:
        ids.add(str(r[id_col]).strip())
        if pos: ids_pos.add(str(r[id_col]).strip())
    for i in loc_cols:
        p = pattern(r[i]); pat[i][p] += 1
        if pos: pat_pos[i][p] += 1
    if args.max_rows and n >= args.max_rows: break
print(f"\nrows: {n:,} | rows with DOC-DOC-LINK=1: {n_pos:,} | distinct patient ids: {len(ids):,} (in linked rows: {len(ids_pos):,})")

print("\ncoverage and format patterns (all rows | linked rows):")
for i in loc_cols:
    tot = sum(pat[i].values()); totp = sum(pat_pos[i].values()) or 1
    cov = 1 - pat[i]["missing"] / tot; covp = 1 - pat_pos[i]["missing"] / totp
    top = ", ".join(f"{k} {v/tot:.0%}" for k, v in pat[i].most_common(4))
    print(f"  {headers[i]:40s} coverage {cov:6.1%} | {covp:6.1%}   patterns: {top}")

# overlap with the prepared bundle's patient set (counts only)
try:
    from referral_gnn.data import default_bundle_path
    pn = pd.read_pickle(os.path.join(default_bundle_path(root), "patient_nodes_PRIVATE.pkl"))
    bundle_ids = set(pn["mso_membership_id"].astype(str))
    print(f"\npatient-id overlap with prepared bundle: bundle {len(bundle_ids):,} | new file {len(ids):,} | intersection {len(bundle_ids & ids):,} | bundle-only {len(bundle_ids - ids):,} | new-only {len(ids - bundle_ids):,}")
except Exception as e:  # noqa: BLE001
    print("bundle overlap check skipped:", type(e).__name__)
